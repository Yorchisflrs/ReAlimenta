"""Genera el Word académico desde la fuente Markdown verificada."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/Requerimientos_No_Funcionales_REAlimenta.md'
TARGET = SOURCE.with_suffix('.docx')
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.05
for name, size in [('Title',24),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
    style = doc.styles[name]
    style.font.name = 'Aptos'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True
    ppr = style._element.get_or_add_pPr()
    border = ppr.find(qn('w:pBdr'))
    if border is not None: ppr.remove(border)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('REAlimenta 2  |  Página ')
r.font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)

def paragraph(text, style=None):
    p = doc.add_paragraph(style=style)
    if ':' in text and text.split(':',1)[0] in ['Categoría','Descripción','Métrica medible','Criterio de aceptación','Procedimiento de prueba','Resultado obtenido','Evidencia disponible']:
        lead, body = text.split(':',1)
        p.add_run(lead+':').bold=True
        p.add_run(body)
    else:
        p.add_run(text)
    return p

def table(lines):
    data = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    data = [row for row in data if not all(re.fullmatch(r'[:\- ]+',c) for c in row)]
    t = doc.add_table(rows=0,cols=len(data[0]))
    t.autofit=False
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=[2.0,2.1,2.9] if len(data[0])==3 else [7/len(data[0])]*len(data[0])
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    for index,record in enumerate(data):
        row=t.add_row()
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index==0: trpr.append(OxmlElement('w:tblHeader'))
        for i,value in enumerate(record):
            cell=row.cells[i]
            cell.width=Inches(widths[i])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=value
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcpr.append(borders)
            margin=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'100');m.set(qn('w:type'),'dxa');margin.append(m)
            tcpr.append(margin)
            if index==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E6EEF5');tcpr.append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.size=Pt(9);r.bold=index==0
    doc.add_paragraph()

lines=SOURCE.read_text(encoding='utf-8').splitlines()
pending=[]
table_lines=[]
first_section=True
def flush():
    if pending:
        paragraph(' '.join(pending))
        pending.clear()

for line in lines+['']:
    if line.startswith('|'):
        flush();table_lines.append(line);continue
    if table_lines: table(table_lines);table_lines=[]
    if not line.strip(): flush();continue
    if line.startswith('# '):
        flush();paragraph(line[2:].title().replace('Realimenta','REAlimenta'),'Title');continue
    if line.startswith('## '):
        flush()
        if first_section:
            doc.add_page_break();first_section=False
        doc.add_heading(line[3:],level=1);continue
    if line.startswith('### '): flush();doc.add_heading(line[4:],level=2);continue
    if line.startswith('- '): flush();paragraph(line[2:],'List Bullet');continue
    if line.endswith('  '): flush();paragraph(line.strip());continue
    pending.append(line.strip())

doc.core_properties.title='Requerimientos no funcionales del sistema REAlimenta'
doc.core_properties.author='Equipo del proyecto REAlimenta'
doc.core_properties.subject='MVP académico y evidencia de verificación'
doc.save(TARGET)
print(TARGET)
